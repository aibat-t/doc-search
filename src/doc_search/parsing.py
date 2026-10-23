"""Этап 1 пайплайна: файл → чистый текст, разбитый на блоки с разделом и страницей."""

import re
from collections import Counter
from dataclasses import dataclass, replace
from pathlib import Path


@dataclass(frozen=True)
class Block:
    """Абзац (или пункт списка) с указанием, в каком разделе и на какой странице он стоит."""

    text: str
    section: str = ""
    page: int | None = None


@dataclass
class ParsedDocument:
    title: str
    blocks: list[Block]


class UnsupportedFormat(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


def parse_file(path: Path, source: str) -> ParsedDocument:
    """source — исходное имя файла: по нему определяем формат и запасной заголовок."""
    suffix = Path(source).suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise UnsupportedFormat(f"Формат {suffix or '(без расширения)'} не поддерживается")
    doc = parser(path)
    blocks = [replace(b, text=clean_text(b.text)) for b in doc.blocks]
    blocks = [b for b in blocks if b.text]
    if not blocks:
        raise EmptyDocument("В документе не найден текст")
    return ParsedDocument(title=doc.title or Path(source).stem, blocks=blocks)


# «Сноска. Статья 1 с изменениями, внесенными Законом РК от ...» (на казахском — «Ескерту.»).
# Удаляем до следующего пункта статьи («2. Текст...») или до конца блока.
_FOOTNOTE = re.compile(
    r"(Сноска|Ескерту)\..*?(?=\s\d+(?:-\d+)?\.\s+[А-ЯЁӘІҢҒҮҰҚӨҺA-Z]|$)", re.DOTALL
)
_SPACE_BEFORE_PUNCT = re.compile(r"\s+([,.;:!?)])")
_SPACE_AFTER_PAREN = re.compile(r"\(\s+")
_MANY_SPACES = re.compile(r"\s{2,}")


def clean_text(text: str) -> str:
    """Чистка, общая для всех форматов: сноски об изменениях и артефакты вроде «слово , слово»."""
    text = _FOOTNOTE.sub("", text)
    text = _SPACE_BEFORE_PUNCT.sub(r"\1", text)
    text = _SPACE_AFTER_PAREN.sub("(", text)
    return _MANY_SPACES.sub(" ", text).strip()


def _section(stack: list[str]) -> str:
    return " > ".join(stack)


def _push(stack: list[str], level: int, heading: str) -> list[str]:
    """Новый заголовок уровня level заменяет всё, что было на этом уровне и глубже."""
    return [*stack[: level - 1], heading]


# ---------- Markdown / текст ----------

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def parse_markdown(text: str) -> ParsedDocument:
    title = ""
    stack: list[str] = []
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block(" ".join(paragraph), _section(stack)))
            paragraph.clear()

    for raw in text.splitlines():
        line = raw.strip()
        heading = _MD_HEADING.match(line)
        if heading:
            flush()
            level, name = len(heading.group(1)), heading.group(2).strip()
            if level == 1 and not title:
                title = name           # «# Заголовок» — название документа, а не раздел
            else:
                stack = _push(stack, max(level - 1, 1), name)
        elif not line:
            flush()                    # пустая строка — граница абзаца
        else:
            paragraph.append(line)
    flush()
    return ParsedDocument(title, blocks)


def _parse_text_file(path: Path) -> ParsedDocument:
    return parse_markdown(path.read_text(encoding="utf-8", errors="replace"))


# ---------- HTML ----------

_HTML_NOISE = ["script", "style", "nav", "footer", "aside", "form", "noscript", "iframe"]
_HTML_TEXT_TAGS = ["p", "li", "pre", "td"]


def parse_html(html: str | bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(_HTML_NOISE):
        tag.decompose()                # меню, футер, скрипты — мусор для поиска

    title = ""
    stack: list[str] = []
    blocks: list[Block] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", *_HTML_TEXT_TAGS]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name == "h1":
            title = title or text
        elif el.name in ("h2", "h3", "h4"):
            stack = _push(stack, int(el.name[1]) - 1, text)
        elif el.find_parent(_HTML_TEXT_TAGS) is None:   # <p> внутри <li> не берём дважды
            blocks.append(Block(text, _section(stack)))

    if not title and soup.title:
        title = soup.title.get_text(strip=True)
    return ParsedDocument(title, blocks)


def _parse_html_file(path: Path) -> ParsedDocument:
    return parse_html(path.read_bytes())


# ---------- DOCX ----------

_DOCX_HEADING = re.compile(r"(heading|заголовок)\s*(\d)")


def _parse_docx_file(path: Path) -> ParsedDocument:
    from docx import Document

    doc = Document(str(path))
    title = ""
    stack: list[str] = []
    blocks: list[Block] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "").lower()
        heading = _DOCX_HEADING.match(style)
        if style == "title":
            title = title or text
        elif heading:
            stack = _push(stack, int(heading.group(2)), text)
        else:
            blocks.append(Block(text, _section(stack)))
    # Таблицы (doc.tables) для простоты пропускаем — в реальном проекте их тоже стоит доставать.
    return ParsedDocument(title, blocks)


# ---------- PDF ----------

_CHAPTER = re.compile(r"^((Глава|Раздел|Тарау|Бөлім)\s+\d+|\d+-(тарау|бөлім))", re.IGNORECASE)
_ARTICLE = re.compile(r"^((Статья|Бап)\s+\d+|\d+-бап)", re.IGNORECASE)


def _normalize(line: str) -> str:
    """Для поиска колонтитулов: «Стр. 3 из 40» и «Стр. 4 из 40» должны считаться одной строкой."""
    return re.sub(r"\d+", "#", line.strip().lower())


def find_boilerplate(pages: list[list[str]], share: float = 0.5, min_pages: int = 3) -> set[str]:
    """Строки, которые повторяются на половине страниц и больше, — колонтитулы и номера страниц."""
    if len(pages) < min_pages:
        return set()
    counts: Counter[str] = Counter()
    for lines in pages:
        counts.update({_normalize(line) for line in lines if line.strip()})
    return {line for line, n in counts.items() if n / len(pages) >= share}


def pdf_lines_to_blocks(pages: list[list[str]]) -> list[Block]:
    """Строки PDF по страницам → блоки. Главы и статьи (законы РК) становятся разделами."""
    boilerplate = find_boilerplate(pages)
    chapter = article = ""
    # Длинное название главы/статьи в PDF переносится на следующую строку:
    #   «Статья 16-1. Уведомление о ... по оказанию услуг»
    #   «связи»
    # Строка со строчной буквы сразу после заголовка — его продолжение, а не текст.
    open_heading: str | None = None      # "chapter" | "article" | None
    blocks: list[Block] = []

    for page_no, lines in enumerate(pages, start=1):
        buffer: list[str] = []

        def flush() -> None:
            if buffer:
                section = " > ".join(s for s in (chapter, article) if s)
                blocks.append(Block(" ".join(buffer), section, page_no))
                buffer.clear()

        for raw in lines:
            line = raw.strip()
            if not line or _normalize(line) in boilerplate:
                continue
            if _CHAPTER.match(line):
                flush()
                chapter, article, open_heading = line, "", "chapter"
            elif _ARTICLE.match(line):
                flush()
                article, open_heading = line, "article"
            elif open_heading and not buffer and line[0].islower():
                if open_heading == "chapter":
                    chapter = f"{chapter} {line}"
                else:
                    article = f"{article} {line}"
            else:
                open_heading = None
                if buffer and buffer[-1].endswith("-"):
                    buffer[-1] = buffer[-1][:-1] + line     # перенос слова: «пере-» + «вод»
                else:
                    buffer.append(line)
        flush()
    return blocks


def _parse_pdf_file(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").splitlines() for page in reader.pages]
    if not any(line.strip() for lines in pages for line in lines):
        raise EmptyDocument("В PDF нет текстового слоя — похоже, это скан, нужен OCR")
    title = (reader.metadata.title if reader.metadata else None) or ""
    return ParsedDocument(title.strip(), pdf_lines_to_blocks(pages))


_PARSERS = {
    ".md": _parse_text_file,
    ".txt": _parse_text_file,
    ".html": _parse_html_file,
    ".htm": _parse_html_file,
    ".docx": _parse_docx_file,
    ".pdf": _parse_pdf_file,
}
SUPPORTED_SUFFIXES = frozenset(_PARSERS)
