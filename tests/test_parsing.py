from pathlib import Path

import pytest

from doc_search.parsing import (
    EmptyDocument,
    UnsupportedFormat,
    find_boilerplate,
    parse_file,
    parse_html,
    parse_markdown,
    pdf_lines_to_blocks,
)


def test_markdown_title_and_sections() -> None:
    doc = parse_markdown("# Правила\n\n## Блокировка\nПервый абзац\nпродолжение.\n\nВторой.\n\n### Сроки\nТретий.")
    assert doc.title == "Правила"
    assert [(b.text, b.section) for b in doc.blocks] == [
        ("Первый абзац продолжение.", "Блокировка"),
        ("Второй.", "Блокировка"),
        ("Третий.", "Блокировка > Сроки"),
    ]


def test_html_drops_navigation_and_keeps_sections() -> None:
    html = """
    <html><head><title>Сайт банка</title></head><body>
      <nav><a>Главная</a><a>Кредиты</a></nav>
      <h1>Тарифы</h1>
      <h2>Переводы</h2><p>Перевод на карту другого банка — 0,95%.</p>
      <ul><li><p>Внутри банка бесплатно.</p></li></ul>
      <footer>© Банк</footer>
    </body></html>
    """
    doc = parse_html(html)
    assert doc.title == "Тарифы"
    assert [b.text for b in doc.blocks] == ["Перевод на карту другого банка — 0,95%.", "Внутри банка бесплатно."]
    assert all(b.section == "Переводы" for b in doc.blocks)


def test_pdf_boilerplate_articles_and_hyphenation() -> None:
    pages = [
        ["Закон о банках", "Глава 1. Общие положения", "Статья 1. Понятия", "Банк — юридиче-", "ское лицо.", "Стр. 1"],
        ["Закон о банках", "Статья 2. Лицензия", "Нужна лицензия.", "Стр. 2"],
        ["Закон о банках", "Глава 2. Отчётность", "Статья 54. Сроки", "Срок — 10 дней.", "Стр. 3"],
    ]
    assert "стр. #" in find_boilerplate(pages)
    blocks = pdf_lines_to_blocks(pages)
    assert [(b.text, b.section, b.page) for b in blocks] == [
        ("Банк — юридическое лицо.", "Глава 1. Общие положения > Статья 1. Понятия", 1),
        ("Нужна лицензия.", "Глава 1. Общие положения > Статья 2. Лицензия", 2),
        ("Срок — 10 дней.", "Глава 2. Отчётность > Статья 54. Сроки", 3),
    ]


def test_kazakh_articles() -> None:
    pages = [["1-тарау. Жалпы ережелер", "1-бап. Негізгі ұғымдар", "Банк — заңды тұлға."]]
    assert pdf_lines_to_blocks(pages)[0].section == "1-тарау. Жалпы ережелер > 1-бап. Негізгі ұғымдар"


def test_docx(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Договор", level=0)          # стиль Title
    document.add_heading("1. Предмет", level=1)
    document.add_paragraph("Банк открывает счёт.")
    path = tmp_path / "contract.docx"
    document.save(str(path))

    doc = parse_file(path, "contract.docx")
    assert doc.title == "Договор"
    assert [(b.text, b.section) for b in doc.blocks] == [("Банк открывает счёт.", "1. Предмет")]


def test_title_falls_back_to_file_name(tmp_path: Path) -> None:
    path = tmp_path / "x.txt"
    path.write_text("Просто текст.", encoding="utf-8")
    assert parse_file(path, "Тарифы 2026.txt").title == "Тарифы 2026"


def test_errors(tmp_path: Path) -> None:
    empty = tmp_path / "empty.md"
    empty.write_text("\n\n", encoding="utf-8")
    with pytest.raises(EmptyDocument):
        parse_file(empty, "empty.md")
    with pytest.raises(UnsupportedFormat):
        parse_file(empty, "table.xlsx")
